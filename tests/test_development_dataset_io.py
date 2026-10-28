"""Small independent acceptance cases for the multiformat development surface."""

from __future__ import annotations

import hashlib
import inspect
import json
import math
import os
import re
import subprocess
import sys
import zipfile
from decimal import Decimal
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from equiframe import cli
from equiframe.datasets import (
    DatasetChoiceRequired,
    DatasetReadError,
    DatasetReadOptions,
    convert_dataset,
    load_dataset,
)


def cli_runner():
    # Typer's native runner echoes visible prompt responses already. Older
    # Typer releases delegate to Click, where echo_stdin is explicit.
    if "echo_stdin" in inspect.signature(CliRunner).parameters:
        return CliRunner(echo_stdin=True)
    return CliRunner()


def make_stdin_interactive(monkeypatch):
    module = inspect.getmodule(CliRunner.isolation)
    monkeypatch.setattr(module._NamedTextIOWrapper, "isatty", lambda self: True)


def write_text(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def simple_file(tmp_path, kind):
    if kind == "json":
        return write_text(
            tmp_path, "source.json", '[{"id":"001","amount":12},{"id":"002","amount":7}]'
        )
    if kind == "jsonl":
        return write_text(
            tmp_path, "source.jsonl", '{"id":"001","amount":12}\n{"id":"002","amount":7}\n'
        )
    if kind in ("csv", "txt"):
        return write_text(tmp_path, "source." + kind, "id,amount\n001,12\n002,7\n")
    if kind == "xml":
        return write_text(
            tmp_path,
            "source.xml",
            "<rows><row><id>001</id><amount>12</amount></row><row><id>002</id><amount>7</amount></row></rows>",
        )
    if kind == "excel":
        path = tmp_path / "source.xlsx"
        book = Workbook()
        book.active.append(["id", "amount"])
        book.active.append(["001", 12])
        book.active.append(["002", 7])
        book.save(path)
        book.close()
        return path
    if kind == "parquet":
        path = tmp_path / "source.parquet"
        pq.write_table(pa.table({"id": ["001", "002"], "amount": [12, 7]}), path)
        return path
    raise AssertionError(kind)


@pytest.mark.parametrize("kind", ["json", "jsonl", "csv", "txt", "xml", "excel", "parquet"])
@pytest.mark.parametrize("target", ["pandas", "polars", "duckdb"])
def test_formats_make_real_native_tables_without_changing_values(tmp_path, kind, target):
    source = simple_file(tmp_path, kind)
    options = DatasetReadOptions(text_mode="delimited") if kind == "txt" else None
    with load_dataset(source, target=target, options=options) as result:
        assert result.info.rows == 2
        assert [column.name for column in result.info.columns] == ["id", "amount"]
        assert isinstance(result.pandas, pd.DataFrame)
        if target == "pandas":
            actual = result.frame.to_dict(orient="records")
        elif target == "polars":
            import polars

            assert isinstance(result.frame, polars.DataFrame)
            actual = result.frame.to_dicts()
        else:
            import duckdb

            assert isinstance(result.frame, duckdb.DuckDBPyRelation)
            actual = result.frame.to_arrow_table().to_pylist()
        assert actual == [{"id": "001", "amount": 12}, {"id": "002", "amount": 7}]


@pytest.mark.parametrize("target", ["pandas", "polars", "duckdb"])
def test_geojson_preserves_geometry_properties_and_collection_metadata(tmp_path, target):
    source = write_text(
        tmp_path,
        "map.geojson",
        json.dumps(
            {
                "type": "FeatureCollection",
                "name": "sample",
                "features": [
                    {
                        "type": "Feature",
                        "id": "001",
                        "properties": {"name": "Quito"},
                        "geometry": {"type": "Point", "coordinates": [-78, 0]},
                    }
                ],
            }
        ),
    )
    with load_dataset(source, target=target) as result:
        row = result.pandas.iloc[0]
        assert row["properties.name"] == "Quito"
        assert row["feature.id"] == "001"
        assert row["geometry"] == {"type": "Point", "coordinates": [-78, 0]}
        assert result.info.metadata["collection_members"]["name"] == "sample"
        native_row = (
            result.frame.to_dict(orient="records")[0]
            if target == "pandas"
            else result.frame.to_dicts()[0]
            if target == "polars"
            else result.frame.to_arrow_table().to_pylist()[0]
        )
        assert native_row["properties.name"] == "Quito"
        assert native_row["feature.id"] == "001"
        assert native_row["geometry"] == {"type": "Point", "coordinates": [-78, 0]}


def test_decimal_integer_and_missing_member_are_not_float_guesses(tmp_path):
    source = write_text(
        tmp_path,
        "exact.json",
        '[{"id":"0007","money":1234567890.123456789,"big":9007199254740993,"missing":null},{"id":"8","money":0.100000000,"big":1}]',
    )
    with load_dataset(source) as result:
        assert result.pandas["money"].tolist() == [
            Decimal("1234567890.123456789"),
            Decimal("0.100000000"),
        ]
        assert result.pandas["big"].tolist() == [9007199254740993, 1]
        assert result.pandas["id"].tolist() == ["0007", "8"]
        assert result.info.metadata["absent_json_members_by_row"] == {"missing": [1]}


@pytest.mark.parametrize("target", ["polars", "duckdb"])
def test_conversion_preserves_arrow_null_nan_and_negative_zero(tmp_path, target):
    source = tmp_path / "special.parquet"
    pq.write_table(pa.table({"value": pa.array([None, float("nan"), -0.0, 2.5])}), source)
    with load_dataset(source, target=target) as result:
        arrow = result.frame.to_arrow() if target == "polars" else result.frame.to_arrow_table()
        values = arrow["value"].to_pylist()
        assert values[0] is None
        assert math.isnan(values[1])
        assert values[2] == 0 and math.copysign(1, values[2]) == -1
        assert values[3] == 2.5


@pytest.mark.parametrize("dimension", [None, "A1:A1"])
def test_excel_reads_actual_cells_when_declared_dimensions_are_wrong(tmp_path, dimension):
    source = simple_file(tmp_path, "excel")
    patched = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, "w") as after:
        for name in before.namelist():
            payload = before.read(name)
            if name == "xl/worksheets/sheet1.xml":
                replacement = (
                    b"" if dimension is None else f'<dimension ref="{dimension}"/>'.encode()
                )
                payload = re.sub(rb"<dimension\b[^>]*/>", replacement, payload)
            after.writestr(name, payload)
    with load_dataset(patched) as result:
        assert result.pandas.to_dict(orient="records") == [
            {"id": "001", "amount": 12},
            {"id": "002", "amount": 7},
        ]


def test_excel_requires_sheet_choice_and_rejects_formulas(tmp_path):
    path = tmp_path / "multiple.xlsx"
    book = Workbook()
    book.active.title = "first"
    book.active.append(["value"])
    book.active.append(["=1+1"])
    second = book.create_sheet("second")
    second.append(["value"])
    second.append([3])
    book.save(path)
    book.close()
    with pytest.raises(DatasetChoiceRequired) as error:
        load_dataset(path)
    assert error.value.option == "sheet"
    assert set(error.value.choices) == {"first", "second"}
    with pytest.raises(DatasetReadError, match=r"formula|Formula"):
        load_dataset(path, options=DatasetReadOptions(sheet="first"))
    with load_dataset(path, options=DatasetReadOptions(sheet="second")) as result:
        assert result.pandas["value"].tolist() == [3]


def test_free_text_keeps_blank_lines_and_does_not_guess_numbers(tmp_path):
    source = write_text(tmp_path, "notes.txt", "001\n\n3.00\n")
    with load_dataset(source, options=DatasetReadOptions(text_mode="lines")) as result:
        assert result.pandas["text"].tolist() == ["001", "", "3.00"]


@pytest.mark.parametrize(
    "name,text",
    [
        ("bad.json", '[{"x":1,"x":2}]'),
        ("bad.json", '[{"x":NaN}]'),
        ("bad.csv", "x,y\n1,2\n3\n"),
        ("bad.xml", '<!DOCTYPE foo [<!ENTITY x "secret">]><rows><row>&x;</row></rows>'),
    ],
)
def test_bad_structure_is_an_error_instead_of_data_loss(tmp_path, name, text):
    with pytest.raises(DatasetReadError):
        load_dataset(write_text(tmp_path, name, text))


def test_explicit_limits_apply_to_actual_rows_and_cells(tmp_path):
    source = simple_file(tmp_path, "excel")
    with pytest.raises(DatasetReadError, match=r"max_rows|max_cells"):
        load_dataset(source, options=DatasetReadOptions(max_rows=1))
    with pytest.raises(DatasetReadError, match=r"max_rows|max_cells"):
        load_dataset(source, options=DatasetReadOptions(max_cells=2))


def test_atomic_export_and_existing_destination(tmp_path):
    source = simple_file(tmp_path, "json")
    destination = tmp_path / "result.parquet"
    with load_dataset(source) as result:
        result.to_parquet(destination)
        assert pq.read_table(destination).to_pylist() == [
            {"id": "001", "amount": 12},
            {"id": "002", "amount": 7},
        ]
        original = destination.read_bytes()
        with pytest.raises(FileExistsError):
            result.to_parquet(destination)
        assert destination.read_bytes() == original
    assert not list(tmp_path.glob(".equiframe-parquet-*.tmp"))


def test_export_write_failure_removes_partial_and_temporary(tmp_path, monkeypatch):
    def denied(table, stream, *args, **kwargs):
        stream.write(b"partial")
        raise PermissionError("injected filesystem write denial")

    source = simple_file(tmp_path, "json")
    destination = tmp_path / "result.parquet"
    monkeypatch.setattr(pq, "write_table", denied)
    with load_dataset(source) as result, pytest.raises(PermissionError):
        result.to_parquet(destination)
    assert not destination.exists()
    assert not list(tmp_path.glob(".equiframe-parquet-*.tmp"))


def test_export_publish_race_does_not_clobber_another_writer(tmp_path, monkeypatch):
    from equiframe import datasets

    source = simple_file(tmp_path, "json")
    destination = tmp_path / "result.parquet"

    def another_writer(temporary, target):
        Path(target).write_bytes(b"other writer")
        raise FileExistsError("injected destination creation race")

    monkeypatch.setattr(datasets.os, "link", another_writer)
    with load_dataset(source) as result, pytest.raises(FileExistsError):
        result.to_parquet(destination)
    assert destination.read_bytes() == b"other writer"
    assert not list(tmp_path.glob(".equiframe-parquet-*.tmp"))


@pytest.mark.parametrize("command", ["load", "inspect"])
def test_cli_success_and_invalid_input_use_separate_channels(tmp_path, command):
    runner = cli_runner()
    source = simple_file(tmp_path, "json")
    flags = ["--target", "polars"] if command == "load" else []
    passed = runner.invoke(cli.app, ["dataset", command, str(source), *flags])
    assert passed.exit_code == 0, passed.output
    assert passed.stderr == ""
    assert '"rows": 2' in passed.stdout
    assert "001" not in passed.stdout
    failed = runner.invoke(cli.app, ["dataset", command, str(tmp_path / "absent.json"), *flags])
    assert failed.exit_code >= 1
    assert failed.stdout == ""
    assert failed.stderr


def test_cli_noninteractive_requires_target_and_handles_eof(tmp_path):
    source = simple_file(tmp_path, "json")
    result = cli_runner().invoke(cli.app, ["dataset", "load", str(source)], input="")
    assert result.exit_code >= 1
    assert result.stdout == ""
    assert "--target" in result.stderr


def test_cli_prompt_selects_engine_in_the_same_invocation(tmp_path, monkeypatch):
    make_stdin_interactive(monkeypatch)
    source = simple_file(tmp_path, "json")
    result = cli_runner().invoke(cli.app, ["dataset", "load", str(source)], input="duckdb\n")
    assert result.exit_code == 0, result.output
    assert '"target_engine": "duckdb"' in result.stdout
    assert '"rows": 2' in result.stderr
    assert '"error"' not in result.stderr


def test_cli_interactive_eof_is_failure_not_success(tmp_path, monkeypatch):
    make_stdin_interactive(monkeypatch)
    source = simple_file(tmp_path, "json")
    result = cli_runner().invoke(cli.app, ["dataset", "load", str(source)], input="")
    assert result.exit_code >= 1
    assert "target_engine" not in result.stdout


@pytest.mark.parametrize("command", ["load", "inspect"])
def test_cli_read_permission_error_is_failure_with_stderr(tmp_path, monkeypatch, command):
    source = simple_file(tmp_path, "json")
    original = Path.open

    def denied(self, *args, **kwargs):
        if self == source:
            raise PermissionError("injected source read denial")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", denied)
    flags = ["--target", "pandas"] if command == "load" else []
    result = CliRunner().invoke(cli.app, ["dataset", command, str(source), *flags])
    assert result.exit_code >= 1
    assert result.stdout == ""
    assert result.stderr


def test_cli_write_permission_error_does_not_emit_success(tmp_path, monkeypatch):
    source = simple_file(tmp_path, "json")
    destination = tmp_path / "result.parquet"

    def denied(*args, **kwargs):
        raise PermissionError("injected output write denial")

    monkeypatch.setattr(pq, "write_table", denied)
    result = CliRunner().invoke(
        cli.app,
        ["dataset", "load", str(source), "--target", "pandas", "--output", str(destination)],
    )
    assert result.exit_code >= 1
    assert result.stdout == ""
    assert result.stderr
    assert not destination.exists()


def test_cli_does_not_require_user_environment_variables(tmp_path, monkeypatch):
    for name in ("EQUIFRAME_CONFIG", "EQUIFRAME_HOME", "R_HOME", "PYTHONPATH"):
        monkeypatch.delenv(name, raising=False)
    source = simple_file(tmp_path, "json")
    result = CliRunner().invoke(cli.app, ["dataset", "load", str(source), "--target", "pandas"])
    assert result.exit_code == 0, result.output
    assert result.stderr == ""


@pytest.mark.parametrize("command", ["load", "inspect"])
def test_native_cli_process_success_and_error_with_minimal_environment(tmp_path, command):
    source = simple_file(tmp_path, "json")
    package_source = Path(__file__).resolve().parents[1] / "src"
    prefix = [
        sys.executable,
        "-I",
        "-B",
        "-c",
        "import sys,runpy;sys.path.insert(0,sys.argv.pop(1));"
        "runpy.run_module('equiframe',run_name='__main__')",
        str(package_source),
    ]
    environment = {
        name: os.environ[name]
        for name in ("SYSTEMROOT", "TEMP", "TMP", "PATH")
        if name in os.environ
    }
    flags = ["--target", "duckdb"] if command == "load" else []
    success = subprocess.run(
        [*prefix, "dataset", command, str(source), *flags],
        capture_output=True,
        text=True,
        input="",
        timeout=30,
        env=environment,
        cwd=tmp_path,
    )
    assert success.returncode == 0, success.stderr
    assert success.stderr == ""
    assert '"rows": 2' in success.stdout
    failure = subprocess.run(
        [*prefix, "dataset", command, str(tmp_path / "missing.json"), *flags],
        capture_output=True,
        text=True,
        input="",
        timeout=30,
        env=environment,
        cwd=tmp_path,
    )
    assert failure.returncode >= 1
    assert failure.stdout == ""
    assert failure.stderr


def test_provenance_identifies_consumed_bytes_effective_options_and_target(tmp_path):
    source = simple_file(tmp_path, "json")
    before = source.read_bytes()
    with load_dataset(source) as loaded:
        provenance = loaded.provenance
        assert provenance.source.sha256 == hashlib.sha256(before).hexdigest()
        assert provenance.source.size_bytes == len(before)
        assert provenance.source.format == "json"
        assert provenance.source.read_options["max_rows"] == 1_000_000
        mutable_copy = provenance.source.read_options
        mutable_copy["max_rows"] = 1
        assert provenance.source.read_options["max_rows"] == 1_000_000
        assert "provenance" not in loaded.info.to_dict()
        assert (
            loaded.info.to_dict(include_provenance=True)["provenance"]["source"]["sha256"]
            == hashlib.sha256(before).hexdigest()
        )
        source.write_text('[{"id":"different","amount":99}]', encoding="utf-8")
        with convert_dataset(loaded, target="polars") as converted:
            assert converted.frame["id"].to_list() == ["001", "002"]
            assert converted.provenance.source == provenance.source
            assert converted.provenance.pandas_sha256 == provenance.pandas_sha256
            assert converted.provenance.target == "polars"
            assert re.fullmatch("[0-9a-f]{64}", converted.provenance.target_sha256)


def test_mutated_pandas_cannot_reuse_earlier_source_identity(tmp_path):
    source = simple_file(tmp_path, "json")
    with load_dataset(source) as loaded:
        loaded.pandas.loc[0, "id"] = "changed"
        with pytest.raises(
            (DatasetReadError, ValueError), match=r"mutat|chang|identit|fingerprint"
        ):
            convert_dataset(loaded, target="polars")


@pytest.mark.skipif(sys.platform != "win32", reason="Real Windows file sharing denial")
@pytest.mark.parametrize("command", ["load", "inspect"])
def test_cli_real_windows_exclusive_file_lock_is_a_read_error(tmp_path, command):
    import ctypes
    from ctypes import wintypes

    source = simple_file(tmp_path, "json")
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    create = kernel.CreateFileW
    create.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    create.restype = wintypes.HANDLE
    close = kernel.CloseHandle
    close.argtypes = [wintypes.HANDLE]
    close.restype = wintypes.BOOL
    handle = create(str(source), 0x80000000, 0, None, 3, 0x80, None)
    assert handle not in (None, wintypes.HANDLE(-1).value), ctypes.get_last_error()
    try:
        flags = ["--target", "pandas"] if command == "load" else []
        result = cli_runner().invoke(cli.app, ["dataset", command, str(source), *flags])
        assert result.exit_code >= 1
        assert result.stdout == ""
        assert result.stderr
        assert "permission" in result.stderr.lower() or "denied" in result.stderr.lower()
    finally:
        assert close(handle)
